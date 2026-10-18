"""Builds the learner-import Excel template.

Every column with a fixed set of answers is a dropdown fed from the standard
lists in helpers.py -- the same lists the import validation checks against --
so what a teacher can pick in Excel is exactly what the system accepts.
Required columns are marked with "*" and a coloured header, and an
Instructions sheet spells out every column's rule.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..helpers import LEARNER_FIELD_OPTIONS

# Rows on the Learners sheet that get dropdowns/validation (row 1 is the header).
TEMPLATE_ROWS = 1000

# (header, field, required, rule shown in the Instructions sheet, example, width)
IMPORT_COLUMNS = [
    ("LRN", "lrn", True, "Exactly 12 digits.", "123456789012", 16),
    ("Last Name", "last_name", True, "Learner's surname.", "Dela Cruz", 18),
    ("First Name", "first_name", True, "Learner's given name.", "Juan", 18),
    ("Middle Name", "middle_name", False, "Leave blank if none.", "Reyes", 16),
    ("Sex", "sex", True, None, "Male", 10),
    ("Date of Birth", "birthdate", True, "A real date, not in the future (YYYY-MM-DD).", "2008-05-14", 15),
    ("Learning Modality", "modality", True, None, "Face-to-Face", 18),
    ("Civil Status", "civil_status", True, None, "Single", 14),
    ("Re-enrollee (Yes/No)", "re_enrollee", False, None, "No", 18),
    ("Employment Status", "employment_status", False, None, "Employed", 18),
    ("Distance from CLC (km)", "distance_from_clc_km", False, "A number, 0 or more.", "5.5", 20),
    ("Contact Number", "contact_number", False, "11 digits starting with 09.", "09171234567", 16),
    ("Guardian Contact Number", "guardian_contact_number", False, "11 digits starting with 09.", "09181234567", 22),
    ("Monthly Income", "monthly_income", False, "Household income in pesos, 0 or more.", "8000", 16),
    ("Socio-economic Status", "occupation", True, None, "Formal Employment", 30),
]

_REQUIRED_FILL = PatternFill("solid", fgColor="FDE2E1")
_OPTIONAL_FILL = PatternFill("solid", fgColor="E8EEF5")
_HEADER_FONT = Font(bold=True, color="1F2937")
_REQUIRED_FONT = Font(bold=True, color="B42318")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_label(header, required):
    return f"{header} *" if required else header


def _add_validation(ws, dv, col_letter):
    dv.add(f"{col_letter}2:{col_letter}{TEMPLATE_ROWS + 1}")
    ws.add_data_validation(dv)


def _build_lists_sheet(wb):
    """Hidden sheet holding each dropdown's allowed values; returns
    {field: absolute range} for the list validations to point at."""
    ws = wb.create_sheet("Lists")
    ranges = {}
    for col, (field, options) in enumerate(LEARNER_FIELD_OPTIONS.items(), start=1):
        letter = get_column_letter(col)
        ws.cell(row=1, column=col, value=field)
        for row, option in enumerate(options, start=2):
            ws.cell(row=row, column=col, value=option)
        ranges[field] = f"Lists!${letter}$2:${letter}${len(options) + 1}"
    ws.sheet_state = "hidden"
    return ranges


def _build_learners_sheet(ws, list_ranges):
    ws.title = "Learners"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for col, (header, field, required, _rule, _example, width) in enumerate(IMPORT_COLUMNS, start=1):
        letter = get_column_letter(col)
        cell = ws.cell(row=1, column=col, value=_header_label(header, required))
        cell.font = _REQUIRED_FONT if required else _HEADER_FONT
        cell.fill = _REQUIRED_FILL if required else _OPTIONAL_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = width

        need = "Required." if required else "Optional."
        dv = None
        if field in list_ranges:
            dv = DataValidation(type="list", formula1=f"={list_ranges[field]}", allow_blank=not required)
            dv.error = f"Please pick a {header} from the dropdown list."
            dv.prompt = f"{need} Pick from the list."
        elif field == "lrn":
            dv = DataValidation(type="textLength", operator="equal", formula1="12")
            dv.error = "LRN must be exactly 12 digits."
            dv.prompt = f"{need} 12 digits."
        elif field in ("contact_number", "guardian_contact_number"):
            dv = DataValidation(type="textLength", operator="equal", formula1="11", allow_blank=True)
            dv.error = "Contact number must be 11 digits starting with 09."
            dv.prompt = f"{need} 11 digits starting with 09."
        elif field == "birthdate":
            dv = DataValidation(
                type="date", operator="between",
                formula1="DATE(1900,1,1)", formula2="TODAY()",
            )
            dv.error = "Please enter a valid date of birth that is not in the future."
            dv.prompt = f"{need} Format: YYYY-MM-DD."
        elif field in ("distance_from_clc_km", "monthly_income"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            dv.error = f"{header} must be a number, 0 or more."
            dv.prompt = f"{need} A number, 0 or more."
        elif required:
            dv = DataValidation(type="textLength", operator="greaterThan", formula1="0")
            dv.error = f"{header} is required."
            dv.prompt = need

        if dv is not None:
            dv.showErrorMessage = True
            dv.showInputMessage = True
            dv.errorTitle = "Invalid value"
            dv.promptTitle = header
            _add_validation(ws, dv, letter)

        # Text format keeps the leading zeros of LRNs and 09xx numbers.
        if field in ("lrn", "contact_number", "guardian_contact_number"):
            for row in range(2, TEMPLATE_ROWS + 2):
                ws.cell(row=row, column=col).number_format = "@"
        elif field == "birthdate":
            for row in range(2, TEMPLATE_ROWS + 2):
                ws.cell(row=row, column=col).number_format = "yyyy-mm-dd"


def _build_instructions_sheet(wb):
    ws = wb.create_sheet("Instructions", 0)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 20

    ws["A1"] = "StayEd Learner Import Template"
    ws["A1"].font = Font(bold=True, size=14)
    notes = [
        "1. Fill in the Learners sheet, one learner per row. Do not rename or reorder the column headers.",
        "2. Columns marked * (red header) are REQUIRED. Rows missing any of them will not be imported.",
        "3. Columns with a dropdown only accept the values in the list. Pick from the dropdown instead of typing.",
        "4. Save as .xlsx and upload it in StayEd. You can review and fix every row before it is imported.",
    ]
    for i, note in enumerate(notes, start=3):
        ws.cell(row=i, column=1, value=note)

    header_row = len(notes) + 4
    for col, label in enumerate(("Column", "Required?", "Allowed values / format", "Example"), start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _OPTIONAL_FILL
        cell.border = _BORDER

    for offset, (header, field, required, rule, example, _w) in enumerate(IMPORT_COLUMNS, start=1):
        row = header_row + offset
        options = LEARNER_FIELD_OPTIONS.get(field)
        allowed = f"Dropdown: {', '.join(options)}" if options else rule
        values = (_header_label(header, required), "Required" if required else "Optional", allowed, example)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if required:
            ws.cell(row=row, column=1).font = _REQUIRED_FONT
            ws.cell(row=row, column=2).font = _REQUIRED_FONT


def build_import_template() -> BytesIO:
    wb = Workbook()
    list_ranges = _build_lists_sheet(wb)
    _build_learners_sheet(wb.active, list_ranges)
    _build_instructions_sheet(wb)
    # Open on the Learners sheet; Instructions is the first tab beside it.
    wb.active = wb.sheetnames.index("Learners")
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out

