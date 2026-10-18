from io import BytesIO

from openpyxl import load_workbook

from app.helpers import (
    CIVIL_STATUS_OPTIONS,
    match_option,
    normalise_contact_number,
)
from app.routes.learner_routes import _check_import_row
from app.services.import_template import IMPORT_COLUMNS, build_import_template


def _row(**overrides):
    row = {
        "lrn": "123456789012",
        "last_name": "Dela Cruz",
        "first_name": "Juan",
        "sex": "Male",
        "birthdate": "2008-05-14",
        "modality": "Face-to-Face",
        "civil_status": "Single",
        "occupation": "Formal Employment",
    }
    row.update(overrides)
    return row


def test_match_option_returns_the_standard_spelling():
    assert match_option("  single ", CIVIL_STATUS_OPTIONS) == "Single"
    assert match_option("MARRIED", CIVIL_STATUS_OPTIONS) == "Married"
    assert match_option("face to face", ("Face-to-Face",)) == "Face-to-Face"
    assert match_option("singel", CIVIL_STATUS_OPTIONS) is None
    assert match_option("", CIVIL_STATUS_OPTIONS) is None


def test_normalise_contact_number_restores_dropped_leading_zero():
    assert normalise_contact_number("9171234567") == "09171234567"
    assert normalise_contact_number("0917-123-4567") == "09171234567"
    assert normalise_contact_number("+639171234567") == "09171234567"
    assert normalise_contact_number("") is None


def test_complete_row_has_no_issues_and_is_standardised():
    clean, issues = _check_import_row(
        _row(sex="female", civil_status="married", modality="modular", re_enrollee="y")
    )
    assert issues == []
    assert clean["sex"] == "Female"
    assert clean["civil_status"] == "Married"
    assert clean["modality"] == "Modular"
    assert clean["re_enrollee"] == "Yes"


def test_civil_status_is_mandatory():
    _, issues = _check_import_row(_row(civil_status=None))
    assert "Civil Status is required" in issues


def test_sex_and_modality_no_longer_default_silently():
    _, issues = _check_import_row(_row(sex="", modality=""))
    assert "Sex is required" in issues
    assert "Learning Modality is required" in issues


def test_value_outside_the_dropdown_is_rejected_but_kept_for_display():
    clean, issues = _check_import_row(_row(civil_status="singel", employment_status="Working"))
    assert any(i.startswith("Civil Status must be one of") for i in issues)
    assert any(i.startswith("Employment Status must be one of") for i in issues)
    assert clean["civil_status"] == "singel"


def test_number_and_contact_format_checks():
    _, issues = _check_import_row(
        _row(contact_number="12345", distance_from_clc_km="far", monthly_income="-5")
    )
    assert "Contact Number must be 11 digits starting with 09" in issues
    assert "Distance from CLC must be a number" in issues
    assert "Monthly Income cannot be negative" in issues


def test_future_birthdate_is_rejected():
    _, issues = _check_import_row(_row(birthdate="2999-01-01"))
    assert "Date of Birth cannot be in the future" in issues


def test_template_marks_required_columns_and_has_dropdowns():
    wb = load_workbook(BytesIO(build_import_template().read()))
    assert wb.sheetnames[0] == "Instructions"
    ws = wb["Learners"]
    headers = [c.value for c in ws[1]]
    assert "Civil Status *" in headers
    assert "Middle Name" in headers
    assert len(headers) == len(IMPORT_COLUMNS)

    civil_col = headers.index("Civil Status *") + 1
    civil_letter = ws.cell(row=1, column=civil_col).column_letter
    list_validations = [
        dv for dv in ws.data_validations.dataValidation
        if dv.type == "list" and f"{civil_letter}2" in str(dv.sqref)
    ]
    assert list_validations, "Civil Status column should be a dropdown"
    assert wb["Lists"].sheet_state == "hidden"


def test_filled_template_is_read_from_the_learners_sheet():
    # The template opens with an Instructions sheet; the importer must read
    # the Learners sheet and ignore the pre-formatted blank rows.
    from app.routes.learner_routes import _canonical_rows, _read_upload

    wb = load_workbook(BytesIO(build_import_template().read()))
    ws = wb["Learners"]
    values = ["123456789012", "Dela Cruz", "Juan", None, "Male", "2008-05-14",
              "Face-to-Face", "Single"]
    for col, value in enumerate(values, start=1):
        ws.cell(row=2, column=col, value=value)
    ws.cell(row=2, column=15, value="Formal Employment")  # Socio-economic Status
    out = BytesIO()
    wb.save(out)

    class Upload:
        filename = "learners.xlsx"

        def read(self):
            return out.getvalue()

    rows = _canonical_rows(_read_upload(Upload()))
    assert len(rows) == 1
    assert rows[0]["civil_status"] == "Single"
    assert _check_import_row(rows[0])[1] == []


def test_socio_economic_status_is_mandatory():
    _, issues = _check_import_row(_row(occupation=None))
    assert "Socio-economic Status is required" in issues
